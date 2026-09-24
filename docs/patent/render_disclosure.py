"""Render the editable patent disclosure and its two schematic figures."""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
STEM = "投机解码反馈_训练数据使用检测_技术交底书"
SOURCE = ROOT / f"{STEM}.md"
OUTPUT = ROOT / f"{STEM}.docx"
FONT_FILE = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_FILE), size)


def box(draw: ImageDraw.ImageDraw, bounds: tuple[int, int, int, int], lines: list[str], *, fill="#f7f8fa") -> None:
    draw.rounded_rectangle(bounds, radius=24, fill=fill, outline="#273142", width=4)
    cx = (bounds[0] + bounds[2]) // 2
    cy = (bounds[1] + bounds[3]) // 2
    gap = 56
    start = cy - gap * (len(lines) - 1) // 2
    for index, line in enumerate(lines):
        draw.text((cx, start + gap * index), line, font=font(38), fill="#172033", anchor="mm")


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
    draw.line((start, end), fill="#273142", width=5)
    x1, y1 = start
    x2, y2 = end
    if x1 == x2:
        tip = [(x2, y2), (x2 - 13, y2 - 21), (x2 + 13, y2 - 21)] if y2 > y1 else [
            (x2, y2), (x2 - 13, y2 + 21), (x2 + 13, y2 + 21)
        ]
    else:
        tip = [(x2, y2), (x2 - 21, y2 - 13), (x2 - 21, y2 + 13)] if x2 > x1 else [
            (x2, y2), (x2 + 21, y2 - 13), (x2 + 21, y2 + 13)
        ]
    draw.polygon(tip, fill="#273142")


def make_figures() -> None:
    canvas = Image.new("RGB", (1800, 930), "white")
    d = ImageDraw.Draw(canvas)
    box(d, (70, 70, 770, 205), ["数据持有人：待检记录"])
    box(d, (1030, 70, 1730, 205), ["可信非成员记录"])
    box(d, (550, 290, 1250, 430), ["候选验证反馈与草稿侧特征"])
    box(d, (1030, 535, 1730, 690), ["非成员参照模型", "及独立校准"])
    box(d, (70, 535, 770, 690), ["位置异常证据", "与记录评分"])
    box(d, (550, 785, 1250, 910), ["检测分数与风险结果"])
    arrow(d, (420, 205), (700, 290))
    arrow(d, (1380, 205), (1100, 290))
    arrow(d, (700, 430), (420, 535))
    arrow(d, (1100, 430), (1380, 535))
    arrow(d, (1030, 612), (770, 612))
    arrow(d, (420, 690), (700, 785))
    arrow(d, (1380, 690), (1100, 785))
    canvas.save(ROOT / "投机解码反馈_方法流程.png", dpi=(220, 220))

    canvas = Image.new("RGB", (1800, 600), "white")
    d = ImageDraw.Draw(canvas)
    positions = [
        (50, 80, 540, 225),
        (655, 80, 1145, 225),
        (1260, 80, 1750, 225),
        (50, 370, 540, 515),
        (655, 370, 1145, 515),
        (1260, 370, 1750, 515),
    ]
    names = ["数据处理模块", "反馈获取模块", "特征提取模块", "非成员参照模型构建模块", "检测评分模块", "结果输出模块"]
    for bounds, name in zip(positions, names):
        box(d, bounds, [name])
    arrow(d, (540, 152), (655, 152))
    arrow(d, (1145, 152), (1260, 152))
    arrow(d, (1505, 225), (1505, 315))
    d.line((1505, 315, 295, 315, 295, 370), fill="#273142", width=5)
    d.polygon([(295, 370), (282, 349), (308, 349)], fill="#273142")
    arrow(d, (540, 442), (655, 442))
    arrow(d, (1145, 442), (1260, 442))
    canvas.save(ROOT / "投机解码反馈_装置结构.png", dpi=(220, 220))


def set_font(run, size: int, *, bold: bool = False) -> None:
    run.font.name = "Noto Serif CJK SC"
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), "Noto Serif CJK SC")


def add_text(doc: Document, line: str, *, kind: str = "body") -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(7 if kind == "body" else 10)
    if kind == "body":
        paragraph.paragraph_format.first_line_indent = Pt(24)
    elif kind == "step":
        paragraph.paragraph_format.keep_with_next = True
    elif kind in {"main", "title", "section"}:
        paragraph.paragraph_format.keep_with_next = True
        if kind in {"main", "title"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(12)
    set_font(paragraph.add_run(line), 16 if kind == "main" else 15 if kind == "title" else 14 if kind == "section" else 12,
             bold=kind in {"main", "title", "section", "step"})


def make_document() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.3), Cm(1.9)
    section.left_margin, section.right_margin = Cm(2.3), Cm(1.9)
    doc.core_properties.title = "一种基于投机解码反馈的训练数据使用检测方法及装置"
    doc.core_properties.subject = "专利技术交底书"

    for line in SOURCE.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("# 说   明   书   摘   要"):
            add_text(doc, "说   明   书   摘   要", kind="main")
        elif line == "# 说   明   书":
            doc.add_page_break()
            add_text(doc, "说   明   书", kind="main")
        elif line.startswith("## "):
            add_text(doc, line[3:], kind="title")
        elif line.startswith("### "):
            add_text(doc, line[4:], kind="section")
        elif line.startswith("步骤（") and re.match(r"步骤（\d+\.）", line):
            add_text(doc, line, kind="step")
        elif line.startswith("!["):
            match = re.match(r"!\[(.+?)\]\((.+?)\)", line)
            if not match:
                raise ValueError(f"invalid image line: {line}")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ROOT / match.group(2)), width=Cm(15.8))
            p.paragraph_format.keep_with_next = True
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_after = Pt(10)
            set_font(caption.add_run(match.group(1)), 10)
        else:
            add_text(doc, line)
    doc.save(OUTPUT)


if __name__ == "__main__":
    make_figures()
    make_document()
    print(OUTPUT)
